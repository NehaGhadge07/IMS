import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_synopsis_doc():
    doc = Document()

    # Set standard 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base font styling
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    def set_para_format(p, space_before=0, space_after=6, line_spacing=1.3, align=WD_ALIGN_PARAGRAPH.LEFT):
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing

    def add_title_heading(text, space_before=12, space_after=18):
        p = doc.add_paragraph()
        set_para_format(p, space_before=space_before, space_after=space_after, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.name = 'Times New Roman'
        return p

    def add_section_heading(text, space_before=12, space_after=10, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = doc.add_paragraph()
        set_para_format(p, space_before=space_before, space_after=space_after, align=align)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = 'Times New Roman'
        return p

    def add_sub_heading(text, space_before=8, space_after=6):
        p = doc.add_paragraph()
        set_para_format(p, space_before=space_before, space_after=space_after)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12.5)
        run.font.name = 'Times New Roman'
        return p

    def add_body_paragraph(text, space_before=0, space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = doc.add_paragraph()
        set_para_format(p, space_before=space_before, space_after=space_after, align=align)
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'
        return p

    def add_bullet(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph(style='List Bullet')
        set_para_format(p, space_before=0, space_after=space_after, align=WD_ALIGN_PARAGRAPH.LEFT)
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.bold = True
            run_b.font.size = Pt(12)
            run_b.font.name = 'Times New Roman'
        run_t = p.add_run(text)
        run_t.font.size = Pt(12)
        run_t.font.name = 'Times New Roman'
        return p

    def add_placeholder_box(title, subtitle):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.0)
        
        # Border and background styling
        shading = parse_xml(r'<w:shd {} w:fill="FAFAFA"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        
        borders = parse_xml(r'''
            <w:tcBorders {} >
                <w:top w:val="dashed" w:sz="8" w:space="0" w:color="999999"/>
                <w:left w:val="dashed" w:sz="8" w:space="0" w:color="999999"/>
                <w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="999999"/>
                <w:right w:val="dashed" w:sz="8" w:space="0" w:color="999999"/>
            </w:tcBorders>
        '''.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(borders)

        p = cell.paragraphs[0]
        set_para_format(p, space_before=80, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
        r1 = p.add_run(f"[ Image Placeholder: {title} ]\n")
        r1.bold = True
        r1.font.size = Pt(13)
        r1.font.name = 'Times New Roman'
        r1.font.color.rgb = RGBColor(80, 80, 80)

        r2 = p.add_run(subtitle + "\n\n(Page left blank according to content for image insertion)")
        r2.italic = True
        r2.font.size = Pt(11)
        r2.font.name = 'Times New Roman'
        r2.font.color.rgb = RGBColor(120, 120, 120)
        
        p_space = cell.add_paragraph()
        set_para_format(p_space, space_before=0, space_after=80, align=WD_ALIGN_PARAGRAPH.CENTER)

    # -------------------------------------------------------------
    # PAGE 1: 1. Introduction
    # -------------------------------------------------------------
    add_title_heading("1.  Introduction", space_before=20, space_after=24)

    add_body_paragraph(
        "The Internship & Talent Management System (IMS) is a comprehensive web-based platform developed to "
        "automate and simplify the internship onboarding, lifecycle tracking, attendance monitoring, task delegation, and "
        "evaluation process in technical training institutions and corporate academies. Developed for IICET Services "
        "Private Limited, the system provides separate, dedicated login and functional modules for the Director (Executive "
        "Management), HR Manager, Technical Trainers (Mentors), and Interns (Students)."
    )
    add_body_paragraph(
        "The Director oversees institutional metrics, financial audits, staff attendance, and executive AI work summaries. "
        "The HR Manager handles student onboarding, batch distribution, fee collection ledgers, leave management, and "
        "document KYC verification. Technical Trainers assign programming tasks, schedule curriculum sessions, and grade "
        "project submissions. Interns securely log in to record daily attendance, submit completed task proofs, track remaining "
        "fee balances, download branded PDF receipts, and access lecture schedules."
    )
    add_body_paragraph(
        "The application is developed using React.js for the frontend, Express.js (Node.js) for the backend REST API, and "
        "SQLite / MySQL as the relational database engine. The system integrates Google Gemini 1.5 Flash AI for automated "
        "weekly/monthly staff work log summarization and employs automated background cron jobs for daily attendance check-out."
    )
    add_body_paragraph(
        "The current internship management process in many training centers and colleges is partially manual or relies on "
        "multiple disconnected tools such as paper attendance sheets, Excel spreadsheets, and informal messaging channels. "
        "Student records, fee payment receipts, and project submissions are frequently lost or unorganized. This conventional "
        "approach is time-consuming, prone to human calculation errors, causes attendance disputes, and delays certification. "
        "Management also lacks unified visibility into daily institutional operations."
    )
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 2: 1.1 Existing System
    # -------------------------------------------------------------
    add_title_heading("1.1  Existing System", space_before=20, space_after=24)

    add_body_paragraph(
        "In many training institutes and colleges, internship operations, onboarding, and student tracking activities are still "
        "managed using manual processes or multiple disconnected systems. Mentors maintain task lists and marks in local "
        "spreadsheets or physical notebooks, making continuous performance evaluation difficult. HR executives record fee payments "
        "and attendance across separate registers, leading to data inconsistencies. Students must repeatedly visit administrative "
        "desks to request fee receipts, verify KYC documents, and track leave approvals. Such legacy systems often lead to:"
    )
    add_bullet("Manual and repetitive data entry across multiple ledgers")
    add_bullet("Time-consuming attendance recording and manual calculation of working hours")
    add_bullet("Human errors in fee calculations, installment records, and payment tracking")
    add_bullet("Difficulty in maintaining centralized student, mentor, and task records")
    add_bullet("Delayed performance evaluation, project feedback, and certificate issuance")
    add_bullet("Lack of centralized document verification and secure KYC storage")
    add_bullet("Absence of executive real-time tracking for management and directors")
    add_bullet("Lack of automated notification and AI-driven performance reporting")
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 3: 1.2 Need and Scope of the New System
    # -------------------------------------------------------------
    add_title_heading("1.2  Need and Scope of the New System", space_before=20, space_after=20)

    add_body_paragraph(
        "The proposed system eliminates manual paperwork and fragmented spreadsheets by providing a single, "
        "centralized, and automated internship management portal."
    )
    add_sub_heading("Need")
    add_bullet("To automate intern onboarding, lifecycle monitoring, and batch allocation.")
    add_bullet("To eliminate manual errors in fee tracking, installment management, and receipt generation.")
    add_bullet("To maintain centralized, normalized database records of interns, mentors, and staff.")
    add_bullet("To automate daily attendance check-ins, check-outs, and late detection through cron schedulers.")
    add_bullet("To provide secure, role-based authentication using JWT tokens and email OTP.")
    add_bullet("To provide automated performance summarization using Google Gemini Generative AI.")
    add_bullet("To enable students and staff to download official fee receipts and audit reports in PDF format.")

    add_sub_heading("Scope")
    add_body_paragraph(
        "The system is designed for professional training organizations, corporate internship programs, and educational "
        "institutions offering technical certifications (such as MCA, BCA, and Engineering). It supports:"
    )
    add_bullet("Centralized Role-Based Access Control (Director, HR, Trainer, Intern)")
    add_bullet("Intern Lifecycle & Batch Management")
    add_bullet("Mentor / Team Lead Allocation & Workload Tracking")
    add_bullet("Multi-Role Attendance Management with Auto-Clockout Cron")
    add_bullet("Task Allocation, Proof Submission, and Mentor Grading Kanban")
    add_bullet("Daily Work Logs & AI-Powered Executive Summarization")
    add_bullet("Course Curriculum Tracks & Live Class Lecture Scheduling")
    add_bullet("Double-Entry Fee Management with Branded PDF Receipt Printing")
    add_bullet("KYC Document Upload & Administrative Approval Repository")
    add_bullet("Leave Requests with Automated Bi-Directional Attendance Sync")
    add_bullet("CRM Admissions Pipeline & Prospective Student Enquiries")
    add_bullet("Internal Multi-Tier Helpdesk Support Ticketing")
    add_bullet("Mock Technical Interviews & Placement Readiness Evaluation")
    add_bullet("Institutional CSV and Tabular Audit Reports")
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 4: 2. Proposed System / 2.1 Objective
    # -------------------------------------------------------------
    add_title_heading("2.  Proposed System", space_before=16, space_after=14)
    add_section_heading("2.1  Objective", space_before=6, space_after=18, align=WD_ALIGN_PARAGRAPH.LEFT)

    add_body_paragraph("The main objectives of the proposed system are:")
    add_bullet("To develop a centralized, normalized web portal for end-to-end internship administration.")
    add_bullet("To provide secure authentication for Director, HR, Trainer, and Intern using JWT and email OTP.")
    add_bullet("To manage intern profiles, batch allocations, and mentor assignments efficiently.")
    add_bullet("To enforce strict attendance tracking with automated late-arrival flags and scheduled cron checkout.")
    add_bullet("To streamline project task allocation, work submission URLs, and mentor review workflows.")
    add_bullet("To manage institutional fee collections, installment balances, and print verifiable PDF receipts.")
    add_bullet("To leverage Google Gemini 1.5 Flash AI to generate concise executive summaries of staff work logs.")
    add_bullet("To synchronize approved leaves automatically with daily attendance records without manual intervention.")
    add_bullet("To provide a structured CRM system to manage admissions inquiries and candidate follow-ups.")
    add_bullet("To improve operational accuracy, ensure 100% data auditability, and eliminate administrative paperwork.")
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 5: 2.2 Hardware and Software Specification
    # -------------------------------------------------------------
    add_title_heading("2.2  Hardware and Software Specification", space_before=16, space_after=20)

    add_sub_heading("Hardware Requirements")
    add_bullet("Processor: Intel Core i3 processor or higher (or AMD Ryzen equivalent)")
    add_bullet("RAM: 8 GB RAM or above")
    add_bullet("Storage: 256 GB SSD or higher")
    add_bullet("Network: Stable Internet Connection (Broadband / Wi-Fi)")
    add_bullet("Display: Standard Monitor with 1366 × 768 resolution or above")

    add_sub_heading("Software Requirements")
    add_bullet("Operating System: Windows 10 / Windows 11 / Linux")
    add_bullet("Programming Environment: Node.js (v20.x or v22.x LTS)")
    add_bullet("Backend Runtime & Framework: Express.js (Node.js REST API)")
    add_bullet("Frontend Library & Build Tool: React.js (v19) with Vite")
    add_bullet("Database Management System: SQLite (Native node:sqlite) / MySQL 8.x")
    add_bullet("Code Editor: Visual Studio Code")
    add_bullet("Version Control: Git & GitHub")
    add_bullet("API Testing Tool: Postman / Thunder Client")
    add_bullet("Web Browser: Google Chrome / Microsoft Edge")
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 6: 2.3 Technology Used
    # -------------------------------------------------------------
    add_title_heading("2.3  Technology Used", space_before=16, space_after=20)

    add_sub_heading("Frontend")
    add_bullet("React.js (v19) — Component-driven Single Page Application framework")
    add_bullet("HTML5 & CSS3 — Responsive, modern user interface design with CSS variables")
    add_bullet("JavaScript (ES6+) — Asynchronous client-side data handling and validation")
    add_bullet("Recharts — Interactive institutional data analytics charts")
    add_bullet("jsPDF & jspdf-autotable — Client-side institutional PDF receipt generator")
    add_bullet("Lucide React — Modern UI icon library")

    add_sub_heading("Backend")
    add_bullet("Node.js — Asynchronous event-driven server runtime environment")
    add_bullet("Express.js — RESTful API routing and modular middleware framework")
    add_bullet("JSON Web Token (JWT) & bcryptjs — Secure token authorization and password hashing")
    add_bullet("Multer — Multipart form data file upload middleware")
    add_bullet("Nodemailer — Automated OTP email delivery via SMTP")
    add_bullet("Google Generative AI (Gemini 1.5 Flash) — AI-powered daily work log summarizer")
    add_bullet("Node-Cron — Automated background attendance and checkout daemon")

    add_sub_heading("Database")
    add_bullet("SQLite (Embedded DatabaseSync) / MySQL 8.x (Normalized 3NF/BCNF Schema)")

    add_sub_heading("Tools & Utilities")
    add_bullet("Visual Studio Code — Source code development environment")
    add_bullet("Git & GitHub — Source code management and repository hosting")
    add_bullet("Postman — REST API endpoint debugging and testing")
    add_bullet("SQLite Inspector / MySQL Workbench — Relational database modeling")
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 7: 2.4 Database Design / ER Diagram
    # -------------------------------------------------------------
    add_title_heading("2.4  Database Design", space_before=16, space_after=14)
    add_section_heading("ER Diagram:", space_before=6, space_after=16, align=WD_ALIGN_PARAGRAPH.LEFT)

    add_placeholder_box(
        "ER Diagram: Entity-Relationship Architecture",
        "Relational Schema for IICET IMS\nEntities: users, mentors, interns, attendance, tasks, task_updates, fees, payments, documents, daily_work_logs, leave_requests, enquiries, company_profile"
    )
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 8: DFD Level 0 (IMAGE INSERTED)
    # -------------------------------------------------------------
    add_section_heading("DFD:", space_before=14, space_after=8)
    add_section_heading("Level 0:", space_before=4, space_after=14)

    dfd0_path = os.path.join(os.path.dirname(__file__), 'synopsis_diagrams', 'dfd_level_0.png')
    if os.path.exists(dfd0_path):
        p_img = doc.add_paragraph()
        set_para_format(p_img, space_before=0, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_img.add_run().add_picture(dfd0_path, width=Inches(6.2))
    else:
        add_placeholder_box("DFD Level 0", "Context Diagram of IMS")
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 9: DFD Level 1: Director (IMAGE INSERTED)
    # -------------------------------------------------------------
    add_section_heading("Level 1:", space_before=14, space_after=8)
    add_section_heading("Director:", space_before=4, space_after=14)

    dfd1_dir_path = os.path.join(os.path.dirname(__file__), 'synopsis_diagrams', 'dfd_level_1_director.png')
    if os.path.exists(dfd1_dir_path):
        p_img = doc.add_paragraph()
        set_para_format(p_img, space_before=0, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_img.add_run().add_picture(dfd1_dir_path, width=Inches(5.6))
    else:
        add_placeholder_box("Level 1: Director", "Director Data Flow Diagram")
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 10: DFD Level 1: HR Manager (IMAGE INSERTED)
    # -------------------------------------------------------------
    add_section_heading("Level 1:", space_before=14, space_after=8)
    add_section_heading("HR Manager:", space_before=4, space_after=14)

    dfd1_hr_path = os.path.join(os.path.dirname(__file__), 'synopsis_diagrams', 'dfd_level_1_hr.png')
    if os.path.exists(dfd1_hr_path):
        p_img = doc.add_paragraph()
        set_para_format(p_img, space_before=0, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_img.add_run().add_picture(dfd1_hr_path, width=Inches(5.6))
    else:
        add_placeholder_box("Level 1: HR Manager", "HR Manager Data Flow Diagram")
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 11: DFD Level 1: Intern (IMAGE INSERTED)
    # -------------------------------------------------------------
    add_section_heading("Level 1:", space_before=14, space_after=8)
    add_section_heading("Intern:", space_before=4, space_after=14)

    dfd1_intern_path = os.path.join(os.path.dirname(__file__), 'synopsis_diagrams', 'dfd_level_1_intern.png')
    if os.path.exists(dfd1_intern_path):
        p_img = doc.add_paragraph()
        set_para_format(p_img, space_before=0, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_img.add_run().add_picture(dfd1_intern_path, width=Inches(5.6))
    else:
        add_placeholder_box("Level 1: Intern", "Intern Data Flow Diagram")
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 12: DFD Level 1: Trainer (PLACEHOLDER AS REQUESTED)
    # -------------------------------------------------------------
    add_section_heading("Level 1:", space_before=14, space_after=8)
    add_section_heading("Trainer / Faculty:", space_before=4, space_after=14)

    add_placeholder_box(
        "Level 1: Trainer (Faculty)",
        "Processes: 1.0 User Authentication | 1.1 View Assigned Batches | 1.2 Manage Tasks & Submissions | 1.3 Conduct Class Sessions | 1.4 Mock Interviews & Tests | 1.5 Submit Daily Work Logs | 1.6 Logout"
    )
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 13: DFD Level 2: Admin / Director (PLACEHOLDER AS REQUESTED)
    # -------------------------------------------------------------
    add_section_heading("Level 2:", space_before=14, space_after=8)
    add_section_heading("Director / Admin:", space_before=4, space_after=14)

    add_placeholder_box(
        "Level 2: Director / Admin Sub-Processes",
        "Detailed decomposition of Director workflows: 1.0 Authentication & RBAC | 2.0 Institutional Analytics | 3.0 Master Staff Management | 4.0 Financial Audits & Fee Controls | 5.0 AI Work Log Verification | 6.0 Institutional Approvals"
    )
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 14: DFD Level 2: HR Manager & Trainer (PLACEHOLDER AS REQUESTED)
    # -------------------------------------------------------------
    add_section_heading("Level 2:", space_before=14, space_after=8)
    add_section_heading("HR Manager & Trainer Sub-Processes:", space_before=4, space_after=14)

    add_placeholder_box(
        "Level 2: HR & Trainer Sub-Processes",
        "Detailed decomposition: HR Onboarding, Fee Collection, Leave Synchronization, Attendance Adjustments, Trainer Task Evaluation, and Mock Interview Scoring"
    )
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 15: 3. Working of Proposed System
    # -------------------------------------------------------------
    add_title_heading("3.  Working of Proposed System", space_before=18, space_after=20)

    add_body_paragraph("The proposed Internship & Talent Management System (IMS) works in the following manner:")
    
    steps = [
        ("1.", "The Director logs into the portal using secure credentials to view institutional analytics, active interns count, financial cash-flow totals, and staff presence."),
        ("2.", "The HR Manager logs in and onboards candidates, establishing their student profile, assigning batch timing, mapping a technical mentor, and generating their fee obligation record."),
        ("3.", "The HR Manager manages CRM prospective inquiries, logs phone follow-ups, and converts prospective candidates directly into enrolled interns."),
        ("4.", "Technical Trainers (Mentors) log into their portal and automatically have their daily arrival logged; they review assigned batches and student rosters."),
        ("5.", "Trainers create structured syllabus tracks and schedule live lecture sessions containing meeting links and downloadable study materials."),
        ("6.", "Trainers create and allocate project tasks to interns with priority levels ('Low', 'Medium', 'High') and submission deadlines."),
        ("7.", "Interns log into their dedicated student portal using their email and password, clock in their daily arrival time (automatically flagged 'Late' if after 09:30 AM), and view assigned tasks."),
        ("8.", "Interns complete assignments and submit live repository / demo links; Trainers review the work, enter marks/feedback, and approve or reject submissions."),
        ("9.", "Trainers conduct technical mock interviews and mock tests, recording quantitative scores (0-10) and feedback for placement readiness."),
        ("10.", "Trainers and HR executives submit daily work logs to the Director detailing accomplishments, hours spent, and blockers."),
        ("11.", "The Director utilizes the integrated Google Gemini 1.5 Flash AI engine to synthesize multi-day work logs into concise executive bullet points."),
        ("12.", "HR records installment payments against tuition balances, and the system instantly generates an official, branded PDF receipt with digital signature lines."),
        ("13.", "Staff and interns submit leave requests; upon approval by HR/Director, the system automatically synchronizes and updates the attendance ledger with 'Leave' status."),
        ("14.", "Background cron jobs at 5:00 PM auto-mark unclocked students as Absent, and at 6:00 PM automatically calculate hours and check out active sessions."),
        ("15.", "The system maintains 100% data auditability, provides downloadable CSV reports, and offers an embedded visual database viewer for system administrators.")
    ]

    for num, txt in steps:
        p_step = doc.add_paragraph()
        set_para_format(p_step, space_before=0, space_after=5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        r_n = p_step.add_run(f"{num} ")
        r_n.bold = True
        r_n.font.size = Pt(11.5)
        r_n.font.name = 'Times New Roman'
        r_t = p_step.add_run(txt)
        r_t.font.size = Pt(11.5)
        r_t.font.name = 'Times New Roman'

    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 16: 4. Advantages
    # -------------------------------------------------------------
    add_title_heading("4.  Advantages", space_before=18, space_after=22)

    advantages = [
        "Automates end-to-end intern onboarding, lifecycle monitoring, and graduation processes.",
        "Reduces manual administrative effort and completely eliminates spreadsheet fragmentation.",
        "Enforces precise daily attendance with automated late-arrival detection and scheduled cron auto-checkout.",
        "Provides strict Role-Based Access Control (RBAC) with tailored views for Director, HR, Trainer, and Intern.",
        "Facilitates seamless task assignment, online proof submission, and mentor evaluation workflows.",
        "Maintains double-entry fee ledgers, prevents overpayments, and generates instant branded PDF receipts.",
        "Integrates Google Gemini 1.5 Flash AI to provide executive performance summaries of staff daily logs.",
        "Implements bi-directional leave-to-attendance auto-synchronization, eliminating attendance discrepancies.",
        "Centralizes candidate KYC document storage with formal administrative approval and rejection workflows.",
        "Embeds a complete CRM lead pipeline for handling student inquiries and follow-ups.",
        "Generates one-click institutional CSV audit reports for attendance, rosters, and financial reconciliations.",
        "Highly scalable architecture built on clean 3rd Normal Form (3NF/BCNF) database standards."
    ]

    for adv in advantages:
        add_bullet(adv, space_after=6)

    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 17: 5. Future Scope
    # -------------------------------------------------------------
    add_title_heading("5.  Future Scope", space_before=18, space_after=22)

    add_body_paragraph("The proposed Internship & Talent Management System can be further enhanced by incorporating:")

    future_scopes = [
        "Native Android & iOS Mobile Application for instant mobile clock-in and push notifications.",
        "Biometric & RFID Attendance Integration for automated hardware-based check-in at facility turnstiles.",
        "AI-Powered Code Review and automated automated testing of intern GitHub project submissions.",
        "Online Payment Gateway Integration (Razorpay / Stripe / UPI QR) for direct student fee payments.",
        "Facial Recognition Attendance verification to eliminate buddy-punching and proxy clock-ins.",
        "Automated Stipend & Payroll Engine linked directly to daily attendance working hours.",
        "Placement Portal Integration connecting certified interns directly with hiring partner corporate recruiters.",
        "SMS & WhatsApp Notification Gateway for instant fee payment receipts and scheduled session reminders.",
        "Cloud Containerization (Docker & Kubernetes) with automated multi-region backup pipelines."
    ]

    for item in future_scopes:
        add_bullet(item, space_after=6)

    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 18: 6. Conclusion
    # -------------------------------------------------------------
    add_title_heading("6.  Conclusion", space_before=20, space_after=24)

    add_body_paragraph(
        "The Internship & Talent Management System (IMS) provides a secure, reliable, and highly efficient institutional "
        "platform for managing technical training cohorts, intern ongoings, attendance records, task evaluations, and financial "
        "fee transactions. It automates repetitive administrative procedures, eliminates human error in fee calculations, enforces "
        "punctuality through automated attendance daemons, and delivers real-time visibility to executive leadership."
    )
    add_body_paragraph(
        "By integrating modern web technologies (React.js, Node.js, and Express) with an impeccably normalized 3NF relational "
        "data model and state-of-the-art Google Gemini 1.5 Flash Generative AI, the platform bridges the gap between students, "
        "instructors, and administration. The system ensures complete transparency, generates verifiable digital PDF receipts, "
        "streamlines KYC compliance, and accelerates institutional workflows. Its modular and extensible architecture ensures it "
        "can readily scale to meet the evolving operational requirements of higher educational institutions and industrial training "
        "academies."
    )
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 19: 7. References
    # -------------------------------------------------------------
    add_title_heading("7.  References", space_before=20, space_after=24)

    references = [
        ("1. React.js Documentation", "https://react.dev/"),
        ("2. Node.js Official Documentation", "https://nodejs.org/docs/"),
        ("3. Express.js Application Framework", "https://expressjs.com/"),
        ("4. SQLite Native Documentation", "https://www.sqlite.org/docs.html"),
        ("5. MySQL 8.0 Reference Manual", "https://dev.mysql.com/doc/refman/8.0/en/"),
        ("6. Google Gemini AI SDK Documentation", "https://ai.google.dev/docs"),
        ("7. JSON Web Token (JWT) Standard", "https://jwt.io/introduction/"),
        ("8. jsPDF Client-side Documentation", "https://github.com/parallax/jsPDF"),
        ("9. Git Version Control System", "https://git-scm.com/doc")
    ]

    for ref_title, ref_url in references:
        p_ref = doc.add_paragraph()
        set_para_format(p_ref, space_before=0, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
        r_t = p_ref.add_run(f"{ref_title} – ")
        r_t.bold = True
        r_t.font.size = Pt(12)
        r_t.font.name = 'Times New Roman'
        r_u = p_ref.add_run(ref_url)
        r_u.font.size = Pt(12)
        r_u.font.name = 'Times New Roman'
        r_u.font.color.rgb = RGBColor(0, 70, 180)

    # Save document
    output_docx_path = os.path.join(os.path.dirname(__file__), 'IMS_Project_Synopsis.docx')
    doc.save(output_docx_path)
    print(f"Successfully generated Word document at: {output_docx_path}")

if __name__ == '__main__':
    create_synopsis_doc()
